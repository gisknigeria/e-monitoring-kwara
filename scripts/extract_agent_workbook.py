"""Read an agent roster without changing the workbook; keep row provenance."""
import json
import re
import sys
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
source, destination = map(Path, sys.argv[1:3])
workbook = openpyxl.load_workbook(source, data_only=True)
records = []
summaries = []
for sheet in workbook:
    merged = {}
    for region in sheet.merged_cells.ranges:
        value = sheet.cell(region.min_row, region.min_col).value
        for row in range(region.min_row, region.max_row + 1):
            for col in range(region.min_col, region.max_col + 1):
                merged[row, col] = value
    count = 0
    wards = set()
    for row in range(1, sheet.max_row + 1):
        def cell(col):
            value = merged.get((row, col), sheet.cell(row, col).value)
            if value is None:
                return ''
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            return re.sub(r'\s+', ' ', str(value)).strip()
        name = cell(2)
        if not name or name.upper() in ['NAME', 'NAMES']:
            continue
        if not re.search(r'[A-Za-z]', name):
            continue
        records.append({'sheet': sheet.title, 'row': row, 'name': name, 'phone': cell(3), 'ward': cell(4), 'pollingUnit': cell(5), 'serial': cell(1)})
        count += 1
        wards.add(cell(4))
    summaries.append({'sheet': sheet.title, 'agents': count, 'wards': sorted(wards)})
destination.parent.mkdir(parents=True, exist_ok=True)
destination.write_text(json.dumps({'source': source.name, 'records': records}, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'total': len(records), 'sheets': summaries}, ensure_ascii=False))
